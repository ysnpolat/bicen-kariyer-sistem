# -*- coding: utf-8 -*-
"""
Rakip web fiyatları (09.10.2026)
--------------------------------
siteler.txt dosyasındaki her sayfayı gerçek bir tarayıcıyla (Chromium) açar,
sayfayı sonuna kadar kaydırıp yüklenen bütün ürünlerin adını ve satış
fiyatını toplar, sonucu Excel'e yazar (sonuclar klasörü).

Ürünler üç yoldan bulunur, hangisi doluysa:
  1. Sayfadaki yapısal veri (schema.org JSON-LD: Product / ItemList)
  2. Mikro veri (itemprop="name" / itemprop="price")
  3. Ürün kartları: fiyat yazısı (12,50 TL / ₺12,50) taşıyan tekrarlı
     kutular; kartın içindeki başlık/bağlantı ürün adı sayılır, üstü çizili
     fiyat "eski fiyat" olarak ayrılır.

Nezaket kuralları: sayfalar sırayla ve aralarında beklenerek açılır; sitenin
robots.txt dosyası otomatik okumayı yasaklıyorsa o sayfa atlanır ve raporda
belirtilir. siteler.txt'de satırın sonuna "| izinli" yazılan sitelerde
(firmanın otomatik okumaya izin verdiği teyit edilmiş) robots.txt
denetlenmez; bu sitelerde sayfalar arası bekleme daha uzundur ve site
erişimi reddederse (403/429) o siteye o çalıştırmada bir daha gidilmez. Reklam izleme parametreleri (gclid, gad_source, utm_…) adresten
otomatik atılır.

Kullanım:  rakip-fiyat-cek.bat            (siteler.txt'deki bütün sayfalar)
           rakip-fiyat-cek.bat <adres>    (tek adresten başlayarak)
           rakip-fiyat-cek.bat --goster   (tarayıcı penceresi ekranda açılır)
"""

import datetime as dt
import os
import pathlib
import re
import sys
import time
import urllib.parse
import urllib.robotparser

KLASOR = os.path.dirname(os.path.abspath(__file__))
SITELER = os.path.join(KLASOR, "siteler.txt")
SONUC = os.path.join(KLASOR, "sonuclar")
KAYDEDILEN = os.path.join(KLASOR, "kaydedilen-sayfalar")

TARAYICI_KIMLIGI = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/126.0 Safari/537.36 BicenFiyatKarsilastirma/1.0"
)
IZLEME = re.compile(r"^(utm_|gad_|gclid|gbraid|wbraid|fbclid|yclid|msclkid|_gl$|mc_)", re.I)
BEKLEME_SN = 3  # sayfalar arası
IZINLI_BEKLEME_SN = 8  # robots.txt denetlenmeyen sitelerde sayfalar arası
EN_AZ_URUN = 150  # bir siteden bu kadar ürün toplanana kadar kategori sayfalarına girilir
EK_SAYFA_SINIRI = 20  # site başına en çok kaç ek sayfa (kategori / sonraki sayfa) açılır
EN_FAZLA_KAYDIRMA = 60


def temiz_adres(adres: str) -> str:
    """Reklam izleme parametrelerini atar."""
    p = urllib.parse.urlsplit(adres.strip())
    q = [(k, v) for k, v in urllib.parse.parse_qsl(p.query, keep_blank_values=True) if not IZLEME.match(k)]
    return urllib.parse.urlunsplit((p.scheme or "https", p.netloc, p.path or "/", urllib.parse.urlencode(q), ""))


def siteleri_oku(argv):
    """(firma, adres, izinli, hedef) listesi. izinli: satır sonunda "| izinli";
    hedef: "| en-az 300" ile değiştirilebilir (varsayılan EN_AZ_URUN)."""
    adresler = [a for a in argv[1:] if a.startswith("http")]
    if adresler:
        return [("", temiz_adres(a), False, EN_AZ_URUN) for a in adresler]

    if not os.path.exists(SITELER):
        return []

    liste = []
    with open(SITELER, encoding="utf-8-sig") as f:
        for satir in f:
            s = satir.strip()
            if not s or s.startswith("#"):
                continue
            # "Firma adı | https://... | izinli" ya da yalnızca adres
            parca = [x.strip() for x in s.split("|")]
            adres = next((x for x in parca if x.startswith("http")), "")
            if not adres:
                continue
            i = parca.index(adres)
            ad = parca[0] if i > 0 else ""
            secenek = [x.lower() for x in parca[i + 1:]]
            izinli = any(x in ("izinli", "izin", "robots-yok") for x in secenek)
            # "en-az 300": bu siteden en az kaç ürün toplanacak
            hedef = EN_AZ_URUN
            for x in secenek:
                m = re.match(r"en[- ]?az\s*:?\s*(\d+)", x)
                if m:
                    hedef = max(1, int(m.group(1)))
            liste.append((ad, temiz_adres(adres), izinli, hedef))
    return liste


_robots = {}


def robots_izin(adres: str) -> bool:
    p = urllib.parse.urlsplit(adres)
    kok = f"{p.scheme}://{p.netloc}"
    if kok not in _robots:
        rp = urllib.robotparser.RobotFileParser()
        rp.set_url(kok + "/robots.txt")
        try:
            rp.read()
        except Exception:
            rp = None  # robots.txt okunamadıysa engel yok sayılır
        _robots[kok] = rp
    rp = _robots[kok]
    return True if rp is None else rp.can_fetch(TARAYICI_KIMLIGI, adres)


def fiyat_sayi(metin):
    """'1.249,90 TL' / '₺12,50' / '12.50' → float"""
    if metin is None:
        return None
    if isinstance(metin, (int, float)):
        return float(metin)
    s = re.sub(r"[^\d.,]", "", str(metin))
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        v = float(s)
        return v if 0 < v < 1_000_000 else None
    except ValueError:
        return None


# Tarayıcı içinde çalışan çıkarıcı -------------------------------------------
CIKARICI = r"""
() => {
  const sonuc = [];
  const ekle = (ad, fiyat, eski, link, kaynak) => {
    ad = (ad || "").replace(/\s+/g, " ").trim();
    if (!ad || fiyat == null || fiyat === "") return;
    sonuc.push({ ad, fiyat: String(fiyat), eski: eski == null ? null : String(eski), link: link || null, kaynak });
  };

  /* 1. JSON-LD */
  const gez = (o) => {
    if (!o || typeof o !== "object") return;
    if (Array.isArray(o)) return o.forEach(gez);
    const tip = [].concat(o["@type"] || []).join(",");
    if (/Product/i.test(tip)) {
      const off = [].concat(o.offers || [])[0] || {};
      const fiyat = off.price ?? off.lowPrice ?? (off.priceSpecification || {}).price;
      ekle(o.name, fiyat, null, o.url || off.url, "json-ld");
    }
    if (o.item) gez(o.item);
    if (o.itemListElement) gez(o.itemListElement);
    if (o["@graph"]) gez(o["@graph"]);
  };
  document.querySelectorAll('script[type="application/ld+json"]').forEach((s) => {
    try { gez(JSON.parse(s.textContent)); } catch (e) {}
  });
  const jsonld = sonuc;

  /* 2. Mikro veri */
  const md = [];
  document.querySelectorAll('[itemtype*="schema.org/Product"]').forEach((k) => {
    const ad = k.querySelector('[itemprop="name"]');
    const f = k.querySelector('[itemprop="price"]');
    const a = k.querySelector("a[href]");
    if (ad && f) md.push({ ad: ad.content || ad.textContent, fiyat: f.getAttribute("content") || f.textContent, eski: null, link: a && a.href, kaynak: "microdata" });
  });
  const mdTemiz = md.map((x) => ({ ...x, ad: x.ad.replace(/\s+/g, " ").trim() }));

  /* 3. Ürün kartları */
  const FIYAT = /(?:₺\s*\d[\d.]*(?:,\d{1,2})?|\d[\d.]*,\d{2}\s*(?:TL|₺)?|\d[\d.]*\s*(?:TL|₺))/i;
  const ustuCizili = (el) => {
    for (let e = el; e && e !== document.body; e = e.parentElement) {
      if (/^(DEL|S|STRIKE)$/.test(e.tagName)) return true;
      const st = getComputedStyle(e);
      if ((st.textDecorationLine || st.textDecoration || "").includes("line-through")) return true;
      if (/old|eski|strike|list-price|was|crossed/i.test(e.className && e.className.baseVal === undefined ? e.className : "")) return true;
    }
    return false;
  };
  const gorunur = (el) => { const r = el.getBoundingClientRect(); return r.width > 0 && r.height > 0; };

  /* Birim fiyat ("11,87 TL / Adet", "45,90 TL/kg", "Birim fiyat"): satış
     fiyatı değil, ayıklanır. Birim bazen yandaki ayrı bir yazıda durduğu
     için kısa ata metnine de bakılır. */
  const BIRIM = /\/\s*(adet|ad\.?|kg|kilo(gram)?|gr|g|lt|l|litre|ml|cl|m|metre|paket|pk|rulo|yaprak|tablet|kaps[uü]l|y[ıi]kama|kullan[ıi]m|100\s*(g|gr|ml))(?![a-zçğıöşü])|birim\s*fiyat|adet\s*fiyat|kg\s*fiyat|litre\s*fiyat/i;
  const birimFiyatMi = (el) => {
    if (BIRIM.test(el.innerText || "")) return true;
    /* Birim, fiyattan hemen sonra ayrı bir yazıda: "11,87 TL" + "/ Adet" */
    const sonra = (((el.nextSibling && el.nextSibling.textContent) || "") + " " +
                   ((el.nextElementSibling && el.nextElementSibling.innerText) || "")).trim();
    if (/^\//.test(sonra) && BIRIM.test(sonra.slice(0, 30))) return true;
    /* Önünde "Birim fiyat" etiketi */
    const once = (((el.previousElementSibling && el.previousElementSibling.innerText) || "") + " " +
                  ((el.previousSibling && el.previousSibling.textContent) || "")).trim();
    return /(birim|adet|kg|litre)\s*fiyat/i.test(once.slice(-30));
  };
  /* "1.299,95 TL" → 1299.95 */
  const sayi = (t) => {
    let m = (String(t || "").match(FIYAT) || [""])[0].replace(/[^\d.,]/g, "");
    if (m.includes(",")) m = m.replace(/\./g, "").replace(",", ".");
    else if (/\.\d{3}$/.test(m)) m = m.replace(/\./g, "");
    const v = parseFloat(m);
    return isNaN(v) ? null : v;
  };
  const yaziBoyu = (el) => parseFloat(getComputedStyle(el).fontSize) || 0;

  /* Fiyat taşıyan en küçük öğeler */
  const fiyatlar = [];
  document.querySelectorAll("body *").forEach((el) => {
    if (["SCRIPT", "STYLE", "NOSCRIPT", "OPTION"].includes(el.tagName)) return;
    const t = (el.innerText || "").trim();
    if (!t || t.length > 40 || !FIYAT.test(t)) return;
    if ([...el.children].some((c) => FIYAT.test((c.innerText || "").trim()))) return;
    if (!gorunur(el)) return;
    if (birimFiyatMi(el)) return;
    fiyatlar.push(el);
  });

  /* Kart: içinde bağlantı + resim ya da başlık olan, en fazla birkaç fiyat
     taşıyan en yakın ata */
  const kartlar = new Map();
  for (const f of fiyatlar) {
    let kart = null;
    for (let e = f.parentElement, d = 0; e && e !== document.body && d < 8; e = e.parentElement, d++) {
      const fiyatSayisi = fiyatlar.filter((x) => e.contains(x)).length;
      if (fiyatSayisi > 4) break;
      if (e.querySelector("a[href]") && (e.querySelector("img") || e.querySelector("h1,h2,h3,h4,h5,[class*=name],[class*=title],[class*=isim],[class*=baslik]"))) {
        kart = e;
        if (e.querySelector("img")) break;
      }
    }
    if (!kart) continue;
    if (!kartlar.has(kart)) kartlar.set(kart, []);
    kartlar.get(kart).push(f);
  }

  /* İç içe kartlarda en dıştaki değil en içteki kalsın */
  const liste = [...kartlar.keys()].filter((k) => ![...kartlar.keys()].some((d) => d !== k && k.contains(d)));

  const urunler = [];
  for (const k of liste) {
    const fl = kartlar.get(k);
    const adEl =
      k.querySelector("[class*=name],[class*=title],[class*=isim],[class*=baslik],[itemprop=name]") ||
      k.querySelector("h1,h2,h3,h4,h5") ||
      [...k.querySelectorAll("a[href]")].find((a) => (a.innerText || a.title || "").trim().length > 3 && !FIYAT.test(a.innerText));
    let ad = adEl ? (adEl.innerText || adEl.getAttribute("title") || "").trim() : "";
    if (!ad) { const img = k.querySelector("img[alt]"); ad = img ? img.alt : ""; }
    ad = ad.split("\n").filter((s) => s.trim() && !FIYAT.test(s)).join(" ");
    const guncel = fl.filter((x) => !ustuCizili(x));
    const eski = fl.filter((x) => ustuCizili(x));

    /* Fiyatlar (12.10.2026): birim fiyatlar zaten ayıklandı.
       - Satış (indirimli) fiyat: üstü çizili olmayanların EN DÜŞÜĞÜ.
         Migros'ta normal fiyat (76,95) düz, "Money ile" fiyat (52,90) ayrı
         bir kutuda yazıyor; büyük yazılanı almak yanlış fiyatı seçiyordu.
       - İndirimsiz fiyat: kartta iki ya da daha çok düz fiyat varsa en
         yükseği (raftaki normal fiyat); yoksa üstü çizili fiyat. Satış
         fiyatının en çok 3 katı (taksit, koli fiyatı gibi uçlar dışarıda). */
    const degerli = fl.map((x) => [x, sayi(x.innerText)]).filter(([, v]) => v != null && v > 0);
    if (!degerli.length) continue;
    const duz = degerli.filter(([x]) => !ustuCizili(x));
    const cizili = degerli.filter(([x]) => ustuCizili(x));
    const satisCifti = (duz.length ? duz : degerli).reduce((a, b) => (b[1] < a[1] ? b : a));
    const secilen = satisCifti[0];
    const sv = satisCifti[1];
    const fiyat = secilen.innerText;

    const enYuksek = (l) =>
      l.filter(([x, v]) => x !== secilen && v > sv + 0.001 && v <= sv * 3).sort((a, b) => b[1] - a[1])[0] || null;
    const eskiCift = duz.length >= 2 ? enYuksek(duz) || enYuksek(cizili) : enYuksek(cizili) || enYuksek(duz);
    const eskiEl = eskiCift ? eskiCift[0] : null;
    const link = (k.querySelector("a[href]") || {}).href;
    urunler.push({ ad, fiyat: fiyat.match(FIYAT)[0], eski: eskiEl ? eskiEl.innerText.match(FIYAT)?.[0] ?? null : null, link, kaynak: "kart" });
  }
  const kart = urunler.filter((u) => u.ad && u.ad.length > 2);

  /* Üç yoldan hangisi en çok ürün bulduysa o; eşitlikte yapısal veri */
  const adaylar = [["json-ld", jsonld], ["microdata", mdTemiz], ["kart", kart]];
  adaylar.sort((a, b) => b[1].length - a[1].length);
  const [yontem, secilenler] = adaylar[0];

  /* Yapısal veride çoğunlukla yalnızca normal fiyat var: aynı ürünün (ad
     ya da bağlantı) kartında indirimli fiyat da görünüyorsa kartın iki
     fiyatı kullanılır (satış = indirimli, eski = normal). */
  if (yontem !== "kart" && kart.length) {
    const anahtar = (t) => String(t || "").toLocaleLowerCase("tr").replace(/\s+/g, " ").trim();
    const kartlarim = new Map();
    for (const u of kart) {
      if (!u.eski) continue;
      kartlarim.set("a:" + anahtar(u.ad), u);
      if (u.link) kartlarim.set("l:" + u.link.split("#")[0], u);
    }
    for (const u of secilenler) {
      const k = (u.link && kartlarim.get("l:" + String(u.link).split("#")[0])) || kartlarim.get("a:" + anahtar(u.ad));
      if (!k) continue;
      u.fiyat = k.fiyat;
      u.eski = k.eski;
    }
  }
  return { yontem, urunler: secilenler };
}
"""


# Aynı sitedeki kategori / liste sayfası bağlantıları ve "sonraki sayfa".
# Ürün sayfaları (ürün kartlarının bağlantıları), hesap, sepet, kurumsal
# sayfalar ve sosyal medya ayıklanır; menüdeki ve adı kategoriye benzeyen
# bağlantılar öne alınır.
BAGLANTILAR = r"""
(urunLinkleri) => {
  const burasi = location.href.split("#")[0];
  const urunler = new Set((urunLinkleri || []).map((x) => String(x).split("#")[0]));
  const HARIC = /(sepet|cart|basket|login|giris|uye-?ol|uyelik|hesab|account|register|kayit|sifre|password|iletisim|contact|hakkimizda|kurumsal|about|kvkk|gizlilik|privacy|cerez|cookie|sozlesme|kosul|kariyer|career|blog|haber|duyuru|magazalar|subeler|store|sss|yardim|help|destek|siparis|order|favori|wishlist|karsilastir|compare|iade|kargo|odeme|checkout|\.(pdf|jpe?g|png|gif|webp|svg|zip)$)/i;
  const KATEGORI = /(kategori|category|categories|\/c\/|-c-\d|reyon|urunler|products|kampanya|firsat|indirim|aktuel|haftanin|market|gida|icecek|temizlik|kozmetik|kisisel|bakim|sut|kahvalti|meyve|sebze|et-|tavuk|balik|atistirmalik|dondurulmus|bebek|ev-|unlu|bakliyat|yag|konserve|sarkuteri|deterjan|kagit)/i;
  const MENU = "nav, header, [class*=menu], [class*=Menu], [class*=categor], [class*=Categor], [class*=kategori], [role=navigation]";

  const puanlar = new Map();
  for (const a of document.querySelectorAll("a[href]")) {
    let u;
    try { u = new URL(a.href, location.href); } catch (e) { continue; }
    if (!/^https?:$/.test(u.protocol) || u.host !== location.host) continue;
    u.hash = "";
    const h = u.href;
    if (h === burasi || urunler.has(h)) continue;
    const yol = decodeURIComponent(u.pathname + u.search).toLowerCase();
    if (yol === "/" || HARIC.test(yol)) continue;
    const metin = (a.innerText || a.title || "").trim();
    if (HARIC.test(metin.toLowerCase())) continue;
    /* Ürün sayfasına benzeyen: kart içinde fiyatın yanındaki bağlantı ya da uzun sayısal kimlik */
    if (/-p-\d+|\/p\/\d+|\/urun\/|\/product\//i.test(yol)) continue;
    let p = 0;
    if (a.closest(MENU)) p += 3;
    if (KATEGORI.test(yol) || KATEGORI.test(metin.toLowerCase())) p += 2;
    if (u.pathname.split("/").filter(Boolean).length > 4) p -= 2;
    if (metin.length === 0) p -= 1;
    puanlar.set(h, Math.max(puanlar.get(h) ?? -99, p));
  }

  const kategoriler = [...puanlar.entries()]
    .filter(([, p]) => p >= 1)
    .sort((a, b) => b[1] - a[1])
    .slice(0, 80)
    .map(([h]) => h);

  /* Sonraki sayfa */
  let sonraki = null;
  const rel = document.querySelector("a[rel=next][href], link[rel=next][href]");
  if (rel) sonraki = rel.href;
  if (!sonraki) {
    for (const a of document.querySelectorAll("a[href]")) {
      const t = (a.innerText || a.getAttribute("aria-label") || a.title || "").trim().toLowerCase();
      if (/^(sonraki( sayfa)?|ileri|next|›|»|>)$/.test(t)) { sonraki = a.href; break; }
    }
  }
  if (sonraki) {
    try {
      const u = new URL(sonraki, location.href);
      sonraki = u.host === location.host && u.href.split("#")[0] !== burasi ? u.href.split("#")[0] : null;
    } catch (e) { sonraki = null; }
  }
  return { kategoriler, sonraki };
}
"""


class EngelHatasi(Exception):
    pass


def sayfa_oku(tarayici, adres):
    sayfa = tarayici.new_page(user_agent=TARAYICI_KIMLIGI, locale="tr-TR", viewport={"width": 1366, "height": 900})
    try:
        yanit = sayfa.goto(adres, wait_until="domcontentloaded", timeout=60_000)
        if yanit is not None and yanit.status in (401, 403, 429):
            # Site istemiyor: zorlanmaz, tekrar denenmez
            raise EngelHatasi(f"site erişimi reddetti (HTTP {yanit.status})")
        try:
            sayfa.wait_for_load_state("networkidle", timeout=15_000)
        except Exception:
            pass

        # Çerez bildirimi gibi katmanları kapatmayı dene
        for metin in ["Kabul Et", "Kabul et", "Tümünü Kabul Et", "Anladım", "Tamam", "Accept"]:
            try:
                d = sayfa.get_by_role("button", name=metin)
                if d.count():
                    d.first.click(timeout=1500)
                    break
            except Exception:
                pass

        # Sonsuz kaydırmalı listeler: yükseklik artmayana kadar kaydır
        onceki = 0
        for _ in range(EN_FAZLA_KAYDIRMA):
            sayfa.mouse.wheel(0, 2500)
            sayfa.wait_for_timeout(700)
            h = sayfa.evaluate("document.body.scrollHeight")
            # "Daha fazla göster" düğmesi
            for metin in ["Daha Fazla", "Daha fazla", "Devamını Gör", "Daha Fazla Ürün"]:
                try:
                    d = sayfa.get_by_role("button", name=re.compile(metin, re.I))
                    if d.count() and d.first.is_visible():
                        d.first.click(timeout=1500)
                        sayfa.wait_for_timeout(1200)
                except Exception:
                    pass
            if h == onceki:
                break
            onceki = h

        sonuc = sayfa.evaluate(CIKARICI)
        baslik = sayfa.title()
        urunler = sonuc.get("urunler", [])
        urun_linkleri = [u.get("link") for u in urunler if u.get("link")]
        try:
            bag = sayfa.evaluate(BAGLANTILAR, urun_linkleri)
        except Exception:
            bag = {"kategoriler": [], "sonraki": None}
        return baslik, sonuc.get("yontem"), urunler, bag
    finally:
        sayfa.close()


def kaydedilenleri_bul():
    """kaydedilen-sayfalar klasöründeki .html / .htm / .mhtml dosyaları"""
    if not os.path.isdir(KAYDEDILEN):
        return []
    return sorted(
        os.path.join(KAYDEDILEN, f)
        for f in os.listdir(KAYDEDILEN)
        if f.lower().endswith((".html", ".htm", ".mhtml", ".mht"))
    )


def kaynak_adres(dosya):
    """Tarayıcının kaydettiği sayfanın asıl adresi: '<!-- saved from url=(..)https://… -->'
    ya da mhtml başlığındaki 'Snapshot-Content-Location'."""
    try:
        with open(dosya, "rb") as f:
            bas = f.read(4000).decode("utf-8", "ignore")
        m = re.search(r"saved from url=\(\d+\)(\S+?)\s*-->", bas) or re.search(r"Snapshot-Content-Location:\s*(\S+)", bas)
        return m.group(1) if m else None
    except Exception:
        return None


def kayitli_sayfa_oku(tarayici, dosya):
    """Kaydedilmiş sayfa: betikler kapalı açılır (sayfa kaydedildiği hâliyle okunur)."""
    baglam = tarayici.new_context(java_script_enabled=False, locale="tr-TR", viewport={"width": 1366, "height": 900})
    try:
        sayfa = baglam.new_page()
        sayfa.goto(pathlib.Path(dosya).resolve().as_uri(), wait_until="load", timeout=60_000)
        sonuc = sayfa.evaluate(CIKARICI)
        return sayfa.title(), sonuc.get("yontem"), sonuc.get("urunler", [])
    finally:
        baglam.close()


def excel_yaz(satirlar, ozet):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    os.makedirs(SONUC, exist_ok=True)
    dosya = os.path.join(SONUC, f"rakip-web-fiyat-{dt.datetime.now():%Y%m%d-%H%M}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Ürünler"
    basliklar = ["Firma", "Ürün adı", "Satış fiyatı", "Eski fiyat", "İndirim %", "Ürün linki", "Sayfa", "Okunma", "Yöntem"]
    ws.append(basliklar)
    for r in satirlar:
        ws.append(r)

    yesil = PatternFill("solid", fgColor="7CB518")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = yesil
        c.alignment = Alignment(vertical="center")
    for satir in ws.iter_rows(min_row=2):
        satir[2].number_format = '#,##0.00 "₺"'
        satir[3].number_format = '#,##0.00 "₺"'
        satir[4].number_format = "0.0"
    for i, gen in enumerate([18, 60, 14, 14, 10, 50, 40, 17, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = gen
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    oz = wb.create_sheet("Özet")
    oz.append(["Firma", "Sayfa", "Durum", "Ürün sayısı", "Yöntem"])
    for r in ozet:
        oz.append(r)
    for c in oz[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = yesil
    for i, gen in enumerate([18, 70, 40, 12, 10], start=1):
        oz.column_dimensions[get_column_letter(i)].width = gen

    wb.save(dosya)
    return dosya


def main():
    siteler = siteleri_oku(sys.argv)
    goster = "--goster" in sys.argv
    tek_sayfa = any(a.startswith("http") for a in sys.argv[1:])
    kaydedilenler = [] if tek_sayfa else kaydedilenleri_bul()

    if not siteler and not kaydedilenler:
        print("siteler.txt içinde sayfa yok ve kaydedilen-sayfalar klasörü boş.")
        print("Her satıra bir adres yazın (isteğe bağlı: Firma | adres) ya da sayfaları tarayıcıdan kaydedin.")
        return 1

    # siteler.txt'deki firma adları: kaydedilmiş sayfanın alan adından firma bulmak için
    firma_adlari = {urllib.parse.urlsplit(a).netloc.replace("www.", ""): f for f, a, _, _ in siteleri_oku([]) if f}

    # Aynı firmanın aynı ürünü (ad + fiyat) birden çok sayfada çıkabiliyor:
    # vitrin, kategori ve kampanya sayfaları. Bir kez yazılır.
    gorulen_firma = {}

    def urunleri_ekle(firma, adres, yontem, urunler):
        gorulen = gorulen_firma.setdefault(firma, set())
        adet = 0
        for u in urunler:
            fiyat = fiyat_sayi(u.get("fiyat"))
            if fiyat is None:
                continue
            eski = fiyat_sayi(u.get("eski"))
            if eski is not None and eski <= fiyat:
                eski = None
            anahtar = (u["ad"].lower(), fiyat)
            if anahtar in gorulen:
                continue
            gorulen.add(anahtar)
            indirim = round((eski - fiyat) / eski * 100, 1) if eski else None
            satirlar.append([firma, u["ad"], fiyat, eski, indirim, u.get("link"), adres, zaman, yontem])
            adet += 1
        return adet

    from playwright.sync_api import sync_playwright

    satirlar, ozet = [], []
    zaman = dt.datetime.now().strftime("%d.%m.%Y %H:%M")

    with sync_playwright() as p:
        tarayici = p.chromium.launch(headless=not goster, **({"executable_path": os.environ["BK_CHROMIUM"]} if os.environ.get("BK_CHROMIUM") else {}))
        try:
            engelli = set()  # bu çalıştırmada erişimi reddeden siteler
            for i, (firma, adres, izinli, hedef) in enumerate(siteler, start=1):
                firma = firma or urllib.parse.urlsplit(adres).netloc.replace("www.", "")
                alan = urllib.parse.urlsplit(adres).netloc
                bekleme = IZINLI_BEKLEME_SN if izinli else BEKLEME_SN
                print(f"[{i}/{len(siteler)}] {firma} · {adres}{'  (izinli)' if izinli else ''} · hedef en az {hedef} ürün")

                if alan in engelli:
                    print("   ATLANDI: site bu çalıştırmada erişimi reddetti.")
                    ozet.append([firma, adres, "Atlandı: site erişimi reddetti", 0, ""])
                    continue

                if not izinli and not robots_izin(adres):
                    print("   ATLANDI: sitenin robots.txt dosyası bu sayfanın otomatik okunmasına izin vermiyor.")
                    print("   Sayfayı tarayıcınızda açıp kaydedin (Ctrl+S, 'Web sayfası, tamamı') ve")
                    print("   kaydedilen-sayfalar klasörüne koyun; bir sonraki çalıştırmada oradan okunur.")
                    print("   Firma otomatik okumaya izin veriyorsa siteler.txt'de satırın sonuna | izinli yazın.")
                    ozet.append([firma, adres, "Atlandı: robots.txt izin vermiyor", 0, ""])
                    continue

                # Verilen sayfa yetmezse (ana sayfada yalnızca vitrin var)
                # aynı sitenin kategori ve "sonraki sayfa" bağlantılarına
                # girilir; hedefe ulaşınca ya da ek sayfa sınırında durulur.
                kuyruk = [adres]
                ziyaret = set()
                toplam = 0
                ek = 0

                while kuyruk:
                    url = kuyruk.pop(0)
                    if url in ziyaret:
                        continue

                    ilk = not ziyaret
                    if not ilk:
                        if toplam >= hedef or ek >= EK_SAYFA_SINIRI:
                            break
                        if not izinli and not robots_izin(url):
                            ziyaret.add(url)
                            continue
                        time.sleep(bekleme)
                        ek += 1

                    ziyaret.add(url)

                    try:
                        baslik, yontem, urunler, bag = sayfa_oku(tarayici, url)
                    except EngelHatasi as h:
                        engelli.add(alan)
                        print(f"   ATLANDI: {h}. Bu siteye bu çalıştırmada bir daha gidilmeyecek.")
                        ozet.append([firma, url, f"Atlandı: {h}", 0, ""])
                        break
                    except Exception as h:
                        print(f"   HATA ({url}): {h}")
                        ozet.append([firma, url, f"Hata: {str(h)[:200]}", 0, ""])
                        if ilk:
                            break
                        continue

                    adet = urunleri_ekle(firma, url, yontem, urunler)
                    toplam += adet
                    if ilk:
                        print(f"   {adet} ürün ({yontem})")
                    else:
                        print(f"   + {adet} yeni ürün · {baslik[:50] or url} (toplam {toplam})")
                    ozet.append([firma, url, f"Okundu · {baslik[:60]}", adet, yontem])

                    # Aynı listenin sonraki sayfası önce, sonra kategoriler
                    if bag.get("sonraki") and adet > 0 and bag["sonraki"] not in ziyaret:
                        kuyruk.insert(0, bag["sonraki"])
                    for k in bag.get("kategoriler") or []:
                        if k not in ziyaret:
                            kuyruk.append(k)

                print(f"   Toplam {toplam} ürün, {len(ziyaret)} sayfa.")
                if toplam < hedef:
                    print(f"   Hedefe ({hedef}) ulaşılamadı: sitede daha fazla kategori bağlantısı bulunamadı ya da")
                    print(f"   {EK_SAYFA_SINIRI} ek sayfa sınırına gelindi. siteler.txt'ye kategori sayfalarını ayrıca ekleyebilirsiniz.")

                if i < len(siteler):
                    time.sleep(bekleme)

            # Elle kaydedilmiş sayfalar: siteye otomatik erişim yok, sayfayı
            # kişi kendi tarayıcısında açıp kaydetmiş.
            for j, dosya in enumerate(kaydedilenler, start=1):
                asil = kaynak_adres(dosya)
                alan = urllib.parse.urlsplit(asil).netloc.replace("www.", "") if asil else ""
                firma = firma_adlari.get(alan) or alan or os.path.splitext(os.path.basename(dosya))[0]
                print(f"[kayıtlı {j}/{len(kaydedilenler)}] {firma} · {os.path.basename(dosya)}")
                try:
                    baslik, yontem, urunler = kayitli_sayfa_oku(tarayici, dosya)
                except Exception as h:
                    print(f"   HATA: {h}")
                    ozet.append([firma, os.path.basename(dosya), f"Hata: {str(h)[:200]}", 0, ""])
                    continue
                adet = urunleri_ekle(firma, asil or os.path.basename(dosya), yontem, urunler)
                print(f"   {adet} ürün ({yontem})")
                ozet.append([firma, asil or os.path.basename(dosya), f"Kaydedilmiş sayfa · {baslik[:60]}", adet, yontem])
        finally:
            tarayici.close()

    dosya = excel_yaz(satirlar, ozet)
    print()
    print(f"Toplam {len(satirlar)} ürün yazıldı:")
    print(f"  {dosya}")
    try:
        if os.name == "nt":
            os.startfile(dosya)  # Excel'de aç
    except Exception:
        pass
    return 0


if __name__ == "__main__":
    sys.exit(main())
